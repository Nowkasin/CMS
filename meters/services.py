# meters/services.py
import pyodbc
from decouple import config
from openpyxl import load_workbook

STATUS_MAP = {
    'สัญญา': '1', 'ต่อสัญญา': '3', 'ไม่ต่อสัญญา': '4',
    'ยกเลิกสัญญา': '5', 'รอปิดสัญญา': '8', 'ปิดสัญญา': '9',
}

COL_KEYS = {
    'YEAR': ['ปีที่สร้างสัญญา'],
    'CUSTOMER_NAME': ['ชื่อลูกหนี้'],
    'TAX_ID': ['เลข13หลัก'],
    'CONTRACT_NO': ['เลขที่สัญญา'],
    'CONTRACT_CODE': ['รหัสสัญญา'],
    'LOCATION': ['Location_id'],       # อาจไม่มีในบางไฟล์ -- ต้องเช็ค None ก่อนใช้
    'LOCATION_NAME': ['สถานที่ตั้ง'],
    'AREA': ['Area_id'],
    'AREA_NAME': ['พื้นที่ตามที่ตั้งพื้นที่เช่า'],
    'SUBAREA': ['SubArea_id'],
    'SUBAREA_NAME': ['พื้นที่ย่อยตามที่ตั้งพื้นที่เช่า'],
    'STATUS': ['สถานะสัญญา'],
    'CONFIRM': ['ยืนยันพื้นที่ที่คิดค่า'],
    'METER_NO': ['เลขประจำเครื่องวัด'],
    'PHASE': ['ระบบไฟฟ้า'],
}

DEFAULT_USER = 'web_upload'


def get_db_connection():
    """
    ต่อ SQL Server ธุรกิจจริงผ่าน pyodbc โดยอ่านค่าเชื่อมต่อจาก .env โดยตรง
    (ไม่ใช่จาก settings.DATABASES['default'] ซึ่งเป็น sqlite ที่ใช้แค่เก็บ session ของ Django เอง)
    ตัวแปรที่ต้องมีใน .env: DB_HOST, DB_PORT, DB_NAME, DB_USER, DB_PASSWORD, DB_ODBC_DRIVER
    """
    driver = config('DB_ODBC_DRIVER', default='ODBC Driver 17 for SQL Server')
    host = config('DB_HOST')
    port = config('DB_PORT', default='1433')
    name = config('DB_NAME')
    user = config('DB_USER')
    password = config('DB_PASSWORD')

    missing = [k for k, v in {
        'DB_HOST': host, 'DB_NAME': name, 'DB_USER': user, 'DB_PASSWORD': password,
    }.items() if not v]
    if missing:
        raise RuntimeError(
            f"ค่าเชื่อมต่อ SQL Server ขาดหายไปใน .env: {', '.join(missing)} "
            f"กรุณาตรวจสอบไฟล์ .env ที่ root โปรเจกต์"
        )

    # escape ปีกกาใน password กัน connection string พังถ้ารหัสผ่านมี { หรือ } อยู่จริง
    safe_password = str(password).replace('}', '}}')

    conn_str = (
        f"DRIVER={{{driver}}};"
        f"SERVER={host},{port};"
        f"DATABASE={name};"
        f"UID={user};"
        f"PWD={{{safe_password}}};"
    )
    return pyodbc.connect(conn_str)


def build_column_map(header_row):
    col_map = {}
    for field, keywords in COL_KEYS.items():
        for i, header in enumerate(header_row):
            if header and all(kw in str(header) for kw in keywords):
                col_map[field] = i
                break
    return col_map


def find_subarea(area_id, subarea_id, location_id=None):
    """
    หาพื้นที่ในฐานข้อมูล คืนค่าเป็น row (Location_id, Area_id, SubArea_id, SubArea_name)
    หรือ None -- เวลาไม่มี location_id ส่งเข้ามา ต้องเอา Location_id จาก row ที่เจอนี้ไปใช้ต่อ
    ตอนบันทึกจริง (sp_Contract_Meter_Save ต้องการ @Location_id แบบ NOT NULL เสมอ)

    ใช้ LTRIM(RTRIM(...)) ทุกฝั่งเทียบ เพราะเจอมาแล้วว่า Contract_hrd_tr.Contract_code
    บางแถวมีอักขระแฝงต่อท้าย (ดู find_contract_by_code) เผื่อ Location_id/Area_id/
    SubArea_id มีปัญหาเดียวกันได้เหมือนกัน กันไว้ก่อนทั้งหมด
    """
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        if location_id:
            cursor.execute(
                """
                SELECT Location_id, Area_id, SubArea_id, SubArea_name
                FROM dbo.Contract_location_subarea_ms
                WHERE LTRIM(RTRIM(Location_id)) = LTRIM(RTRIM(?))
                  AND LTRIM(RTRIM(Area_id)) = LTRIM(RTRIM(?))
                  AND LTRIM(RTRIM(SubArea_id)) = LTRIM(RTRIM(?))
                """,
                location_id, area_id, subarea_id,
            )
        else:
            cursor.execute(
                """
                SELECT Location_id, Area_id, SubArea_id, SubArea_name
                FROM dbo.Contract_location_subarea_ms
                WHERE LTRIM(RTRIM(Area_id)) = LTRIM(RTRIM(?))
                  AND LTRIM(RTRIM(SubArea_id)) = LTRIM(RTRIM(?))
                """,
                area_id, subarea_id,
            )
        return cursor.fetchone()
    finally:
        conn.close()


def find_contract_by_code(contract_code):
    """
    เช็คว่ารหัสสัญญานี้มีอยู่จริงใน Contract_hrd_tr หรือไม่

    ใช้ TRIM ทั้ง 2 ฝั่งก่อนเทียบ เพราะเจอแล้วว่าบางแถวใน Contract_hrd_tr
    มีอักขระที่มองไม่เห็น (ช่องว่างเกิน/ขึ้นบรรทัดใหม่) ติดอยู่ท้ายค่า Contract_code
    (เช่น 'AI5-AI6/2569' จริงๆ เก็บเป็น 'AI5-AI6/2569 ' ในฐานข้อมูล) ทำให้เทียบแบบ
    exact match ตรงๆ ไม่เจอ ทั้งที่ข้อมูลมีอยู่จริง
    """
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT Contract_id, Contract_code FROM dbo.Contract_hrd_tr "
            "WHERE LTRIM(RTRIM(Contract_code)) = LTRIM(RTRIM(?))",
            contract_code,
        )
        return cursor.fetchone()
    finally:
        conn.close()


def find_contract_sheet(sheetnames, keyword):
    candidates = [n for n in sheetnames if n.startswith('สัญญาปี') and keyword in n]
    if len(candidates) > 1:
        raise ValueError(
            f"พบชีทสัญญาที่ขึ้นต้นด้วย 'สัญญาปี' และมีคำว่า {keyword!r} มากกว่า 1 ชีท: {candidates} "
            f"กรุณาตรวจสอบชื่อชีทในไฟล์ Excel"
        )
    return candidates[0] if candidates else None


def parse_excel_staged(uploaded_file):
    wb = load_workbook(uploaded_file, data_only=True)
    water_sheet = find_contract_sheet(wb.sheetnames, 'น้ำ')
    elec_sheet = find_contract_sheet(wb.sheetnames, 'ไฟ')

    rows_out = []
    counter = {'idx': 0}

    def handle_sheet(sheet_name, type_cd):
        if not sheet_name:
            return
        type_label = 'น้ำ' if type_cd == 9 else 'ไฟฟ้า'

        ws = wb[sheet_name]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        col = build_column_map(header_row)

        required = ['SUBAREA', 'AREA', 'CONTRACT_NO', 'CONTRACT_CODE', 'CONFIRM', 'METER_NO']
        missing = [f for f in required if f not in col]
        if missing:
            raise ValueError(
                f"ชีท '{sheet_name}' หาคอลัมน์ไม่เจอ: {missing} "
                f"กรุณาตรวจสอบชื่อ header แถวแรกของชีทนี้"
            )
        has_location = 'LOCATION' in col

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue

            subarea_id = row[col['SUBAREA']]
            area_id = row[col['AREA']]
            loc_id = row[col['LOCATION']] if has_location else None
            if not subarea_id or not area_id:
                continue
            area_id, subarea_id = str(area_id).strip(), str(subarea_id).strip()
            loc_id = str(loc_id).strip() if loc_id else None

            contract_id = row[col['CONTRACT_NO']]
            contract_code = row[col['CONTRACT_CODE']]
            contract_code = str(contract_code).strip() if contract_code else contract_code
            confirm_flag = row[col['CONFIRM']]
            meter_no_raw = row[col['METER_NO']]

            counter['idx'] += 1
            entry = {
                'idx': counter['idx'],
                'type_cd': type_cd, 'type_label': type_label,
                'contract_year': row[col['YEAR']] if 'YEAR' in col else None,
                'customer_name': row[col['CUSTOMER_NAME']] if 'CUSTOMER_NAME' in col else None,
                'tax_id': row[col['TAX_ID']] if 'TAX_ID' in col else None,
                'location_id': loc_id, 'area_id': area_id,
                'subarea_id': subarea_id,
                'subarea_name': row[col['SUBAREA_NAME']] if 'SUBAREA_NAME' in col else None,
                'contract_id': contract_id, 'contract_code': contract_code,
                'meter_no': None, 'meter_no_status': None, 'phase_type': None,
                'status': 'ok', 'message': '',
            }

            if confirm_flag != 1 or meter_no_raw == 'ไม่มีมิเตอร์':
                entry['status'] = 'skip'
                entry['message'] = f'ไม่มีการคิดค่า{type_label}ในพื้นที่นี้'
                rows_out.append(entry)
                continue

            # --- เช็คพื้นที่ + แก้บั๊ก Location_id เป็น None ---
            subarea_match = find_subarea(area_id, subarea_id, loc_id)
            if not subarea_match:
                where = f'{loc_id}/{area_id}/{subarea_id}' if loc_id else f'{area_id}/{subarea_id}'
                entry['status'] = 'error'
                entry['message'] = f'ไม่พบพื้นที่ {where} ในฐานข้อมูล'
                rows_out.append(entry)
                continue

            if not loc_id:
                loc_id = str(subarea_match[0])
                entry['location_id'] = loc_id

            # อ่านเลขมิเตอร์/เฟสจากไฟล์ก่อนเสมอ ไม่ว่าจะเช็ครหัสสัญญาผ่านหรือไม่
            # (แถวไหน error ก็ยังเห็นข้อมูลจริงจากไฟล์ในตาราง review ได้ ไม่ใช่ว่างเปล่า)
            meter_no = None if meter_no_raw in (None, 'มิเตอร์ไม่มีเลข (GEN เลข)') else str(meter_no_raw)
            entry['meter_no'] = meter_no
            entry['meter_no_status'] = 'รอ Gen เลข' if meter_no is None else 'ปกติ'
            if type_cd == 8 and 'PHASE' in col and row[col['PHASE']]:
                entry['phase_type'] = '3เฟส' if '3' in str(row[col['PHASE']]) else '1เฟส'

            # --- เช็ครหัสสัญญาต้องมีอยู่จริงในระบบ (ของเดิมเช็คแค่ไม่ว่างเปล่า) ---
            if not contract_code:
                entry['status'] = 'error'
                entry['message'] = 'ไม่พบรหัสสัญญาในแถวนี้'
                rows_out.append(entry)
                continue

            contract_match = find_contract_by_code(contract_code)
            if not contract_match:
                entry['status'] = 'error'
                entry['message'] = f"ไม่พบรหัสสัญญา '{contract_code}' ในระบบ"
                rows_out.append(entry)
                continue

            # ใช้ Contract_id จริงจาก DB (ตัวเลข ไม่มีปัญหาเรื่องช่องว่าง/อักขระแฝงแบบข้อความ)
            # แทน contract_code ตอนผูกกับสัญญาจริงใน commit_staged_rows -- กันปัญหาเดิม
            # ซ้ำอีกรอบตอน SP เทียบ Contract_code แบบ exact match ภายใน
            entry['db_contract_id'] = contract_match[0]

            rows_out.append(entry)

    handle_sheet(water_sheet, 9)
    handle_sheet(elec_sheet, 8)

    summary = {
        'total': len(rows_out),
        'ok': sum(1 for r in rows_out if r['status'] == 'ok'),
        'skip': sum(1 for r in rows_out if r['status'] == 'skip'),
        'error': sum(1 for r in rows_out if r['status'] == 'error'),
    }

    return {
        'rows': rows_out,
        'summary': summary,
        'water_sheet_found': bool(water_sheet),
        'elec_sheet_found': bool(elec_sheet),
    }


def sp_meter_save(cursor, location_id, area_id, subarea_id, type_cd, meter_no,
                   meter_no_status, user_id, logs, phase_type=None, meter_id=None):
    """
    meter_id=None (ค่าเริ่มต้น) -> SP จะ INSERT มิเตอร์ใหม่ (ใช้ตอน import จากไฟล์ Excel)
    meter_id=<เลขจริง> -> SP จะ UPDATE มิเตอร์ตัวนั้นแทน (ใช้ตอนแก้ไขจากหน้า dashboard)
    """
    cursor.execute(
        """
        EXEC dbo.sp_Contract_Meter_Save
            @Meter_id        = ?,
            @Location_id     = ?,
            @Area_id         = ?,
            @SubArea_id      = ?,
            @Meter_type_cd   = ?,
            @Meter_no        = ?,
            @Meter_no_status = ?,
            @Phase_type      = ?,
            @UserId          = ?
        """,
        meter_id, location_id, area_id, subarea_id, type_cd, meter_no, meter_no_status, phase_type, user_id
    )
    row = cursor.fetchone()
    result_meter_id = int(row[0]) if row else meter_id
    cls = 'water' if type_cd == 9 else 'electric'
    action = 'UPDATE' if meter_id else 'INSERT'
    logs.append((cls, f"EXEC sp_Contract_Meter_Save ({action})  @SubArea_id='{subarea_id}', "
                       f"@Meter_type_cd={type_cd}, @Meter_no={meter_no!r}  -> Meter_id={result_meter_id}"))
    return result_meter_id


def sp_bind_to_contract(cursor, contract_id, location_id, area_id, subarea_id, meter_id, user_id, logs):
    """
    ใช้ @Contract_id (ตัวเลข) แทน @Contract_code (ข้อความ) เพราะพบว่าบาง Contract_code
    ในฐานข้อมูลมีช่องว่าง/อักขระแฝงต่อท้าย ทำให้ SP เทียบแบบ exact match ไม่เจอ
    ถึงแม้จะ TRIM ตรวจสอบผ่านตอน parse แล้วก็ตาม -- ตัวเลขไม่มีปัญหานี้
    """
    cursor.execute(
        """
        EXEC dbo.sp_Contract_Meter_BindToContract
            @Contract_id   = ?,
            @Location_id   = ?,
            @Area_id       = ?,
            @SubArea_id    = ?,
            @Meter_id_list = ?,
            @UserId        = ?
        """,
        contract_id, location_id, area_id, subarea_id, str(meter_id), user_id
    )
    logs.append(('bind', f"EXEC sp_Contract_Meter_BindToContract  @Contract_id={contract_id}, "
                          f"@Meter_id_list='{meter_id}'"))


def commit_staged_rows(rows, user_id):
    logs = []
    stats = {'meter_ok': 0, 'skipped': 0, 'errors': 0}
    result_rows = []

    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        for r in rows:
            row_result = dict(r)

            if r['status'] == 'skip':
                logs.append(('skip', f"SKIP  SubArea={r['subarea_id']} สัญญา {r['contract_code'] or '-'} -- {r['message']}"))
                stats['skipped'] += 1
                row_result['final_status'] = 'skip'
                result_rows.append(row_result)
                continue

            if r['status'] == 'error':
                logs.append(('err', f"ERROR  SubArea={r['subarea_id']} -- {r['message']}"))
                stats['errors'] += 1
                row_result['final_status'] = 'error'
                result_rows.append(row_result)
                continue

            try:
                meter_id = sp_meter_save(
                    cursor, r['location_id'], r['area_id'], r['subarea_id'], r['type_cd'],
                    r['meter_no'], r['meter_no_status'], user_id, logs, r.get('phase_type'),
                )
                stats['meter_ok'] += 1
                row_result['meter_id'] = meter_id
                row_result['final_status'] = 'success'

                if r.get('db_contract_id') and meter_id:
                    sp_bind_to_contract(
                        cursor, r['db_contract_id'], r['location_id'], r['area_id'],
                        r['subarea_id'], meter_id, user_id, logs,
                    )
            except Exception as exc:
                logs.append(('err', f"ERROR  SubArea={r['subarea_id']} -- {exc}"))
                stats['errors'] += 1
                row_result['final_status'] = 'error'
                row_result['message'] = str(exc)

            result_rows.append(row_result)

        conn.commit()
    finally:
        conn.close()

    return {'logs': logs, 'stats': stats, 'rows': result_rows}


def _to_sql_like_pattern(search):
    """
    แปลงคำค้นแบบที่คนคุ้นเคย (Excel/Windows) เป็น SQL LIKE pattern:
      ?  -> ตัวอักษรใดก็ได้ 1 ตัว   (SQL: _)
      *  -> ตัวอักษรใดก็ได้ 0 ตัวขึ้นไป (SQL: %)
    ต้อง escape % และ _ ตัวจริงที่ผู้ใช้พิมพ์มาก่อน ไม่งั้นถ้าเลขสัญญามี % หรือ _ อยู่จริง
    จะถูกตีความเป็น wildcard ของ SQL ไปโดยไม่ตั้งใจ
    """
    escaped = search.replace('[', '[[]').replace('%', '[%]').replace('_', '[_]')
    escaped = escaped.replace('?', '_').replace('*', '%')
    return f"%{escaped}%"


def fetch_subareas(search=None):
    """
    search: ถ้าระบุ -- กรองเฉพาะแถวที่ contract_code ตรงกับคำค้น
            รองรับ wildcard แบบ Excel: ? = 1 ตัวอักษรใดก็ได้, * = กี่ตัวก็ได้
    """
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        sql = """
            SELECT
                m.Location_id,
                m.Area_id,
                m.SubArea_id,
                s.SubArea_name,
                MAX(CASE WHEN m.Meter_type_cd = 9 THEN m.Meter_no END) AS water_meter_no,
                MAX(CASE WHEN m.Meter_type_cd = 8 THEN m.Meter_no END) AS electric_meter_no,
                MAX(c.Contract_code) AS contract_code,
                MAX(cu.Customer_id) AS customer_id,
                MAX(cu.CompanyName) AS customer_name
            FROM Contract_meter_ms m
            LEFT JOIN Contract_location_subarea_ms s ON s.SubArea_id = m.SubArea_id
            LEFT JOIN Contract_meter_tr t ON t.Meter_id = m.Meter_id AND t.UseOrNot = 1
            LEFT JOIN Contract_hrd_tr c ON c.Contract_id = t.Contract_id
            LEFT JOIN Contract_customer_tr cu ON cu.Contract_id = c.Contract_id
            GROUP BY m.Location_id, m.Area_id, m.SubArea_id, s.SubArea_name
        """
        params = []
        if search:
            sql += " HAVING MAX(c.Contract_code) LIKE ? ESCAPE '['"
            params.append(_to_sql_like_pattern(search))
        sql += " ORDER BY m.Location_id, m.Area_id, m.SubArea_id"

        cursor.execute(sql, params)
        columns = [c[0] for c in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        conn.close()


def fetch_contract_meter_detail(contract_id=None, contract_code=None):
    if contract_id is None and contract_code is None:
        raise ValueError('ต้องระบุ contract_id หรือ contract_code อย่างใดอย่างหนึ่ง')

    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        cursor.execute(
            "EXEC dbo.sp_Contract_Meter_SelectByContract @Contract_id = ?, @Contract_code = ?",
            contract_id, contract_code,
        )
        columns = [c[0] for c in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    finally:
        conn.close()


def fetch_subarea_meters(subarea_id):
    """
    ดึงข้อมูลพื้นที่ + มิเตอร์น้ำ/ไฟที่มีอยู่ (ถ้ามี) ของ SubArea นี้ ใช้เปิดหน้าแก้ไขจาก dashboard
    คืนค่า None ถ้าไม่พบ SubArea นี้เลย
    """
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        cursor.execute(
            """
            SELECT s.Location_id, s.Area_id, s.SubArea_id, s.SubArea_name, a.Area_name
            FROM dbo.Contract_location_subarea_ms s
            LEFT JOIN dbo.Contract_location_area_ms a ON a.Area_id = s.Area_id
            WHERE s.SubArea_id = ?
            """,
            subarea_id,
        )
        subarea_row = cursor.fetchone()
        if not subarea_row:
            return None
        columns = [c[0] for c in cursor.description]
        result = dict(zip(columns, subarea_row))

        cursor.execute(
            """
            SELECT Meter_id, Meter_type_cd, Meter_no, Meter_no_status, Phase_type
            FROM dbo.Contract_meter_ms
            WHERE SubArea_id = ? AND UseOrNot = 1
            """,
            subarea_id,
        )
        columns2 = [c[0] for c in cursor.description]
        meters = [dict(zip(columns2, r)) for r in cursor.fetchall()]
        result['water'] = next((m for m in meters if m['Meter_type_cd'] == 9), None)
        result['electric'] = next((m for m in meters if m['Meter_type_cd'] == 8), None)

        # ยืนยันจาก schema จริงแล้ว: ตาราง Contract_hrd_tr ไม่มีคอลัมน์ "เลขที่สัญญา" แยกต่างหาก
        # (มีแค่ Contract_id ซึ่งเป็น PK, int กับ Contract_code) -- "เลขที่สัญญา" ที่ผู้ใช้กรอกใน
        # Excel ตอน import (ดู parse_excel_staged/CONTRACT_NO) หมายถึง Contract_id ตัวนี้เอง
        cursor.execute(
            """
            SELECT TOP 1 c.Contract_code, c.Contract_id, cu.Customer_id, cu.CompanyName
            FROM dbo.Contract_meter_tr t
            JOIN dbo.Contract_hrd_tr c ON c.Contract_id = t.Contract_id
            LEFT JOIN dbo.Contract_customer_tr cu ON cu.Contract_id = c.Contract_id
            WHERE t.SubArea_id = ? AND t.UseOrNot = 1
            """,
            subarea_id,
        )
        contract_row = cursor.fetchone()
        if contract_row:
            result['contract_code'] = contract_row[0]
            result['contract_no'] = contract_row[1]
            result['customer_id'] = contract_row[2]
            result['customer_name'] = contract_row[3]
        else:
            result['contract_code'] = None
            result['contract_no'] = None
            result['customer_id'] = None
            result['customer_name'] = None

        return result
    finally:
        conn.close()


def save_subarea_meters(location_id, area_id, subarea_id, user_id,
                         water_enabled, water_meter_id, water_meter_no,
                         electric_enabled, electric_meter_id, electric_meter_no, electric_phase):
    """
    บันทึกมิเตอร์น้ำ/ไฟของ SubArea นี้จากหน้าแก้ไขใน dashboard -- update ทันที (ไม่มีหน้ายืนยันเพิ่ม)
    ถ้ามี *_meter_id อยู่แล้ว (เคยมีมิเตอร์ประเภทนี้อยู่ก่อน) จะเป็นการ UPDATE ตัวเดิม
    ถ้ายังไม่มี (*_meter_id เป็น None) จะเป็นการ INSERT มิเตอร์ใหม่ให้พื้นที่นี้

    คืนค่า (logs, result_meter_ids) โดย result_meter_ids = {'water': <meter_id หรือ None>, 'electric': <meter_id หรือ None>}
    -- ต้องคืน meter_id กลับไปด้วย เพราะถ้าเพิ่งมีการ INSERT มิเตอร์ใหม่ (ไม่เคยมี water_meter_id/electric_meter_id
    มาก่อน) ผู้เรียกจะไม่รู้เลขนั้นเลยถ้าไม่ส่งกลับมา แล้วจะเอาไปบันทึกเลขอ่านมิเตอร์ต่อให้ถูกตัวไม่ได้
    """
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        logs = []
        result_meter_ids = {'water': None, 'electric': None}

        if water_enabled:
            meter_no_status = 'รอ Gen เลข' if not water_meter_no else 'ปกติ'
            result_meter_ids['water'] = sp_meter_save(
                cursor, location_id, area_id, subarea_id, 9, water_meter_no or None,
                meter_no_status, user_id, logs, phase_type=None, meter_id=water_meter_id,
            )

        if electric_enabled:
            meter_no_status = 'รอ Gen เลข' if not electric_meter_no else 'ปกติ'
            result_meter_ids['electric'] = sp_meter_save(
                cursor, location_id, area_id, subarea_id, 8, electric_meter_no or None,
                meter_no_status, user_id, logs, phase_type=electric_phase, meter_id=electric_meter_id,
            )

        conn.commit()
        return logs, result_meter_ids
    finally:
        conn.close()


def fetch_readings_for_meters(meter_ids, billing_year_be, billing_month):
    """
    ดึงเลขอ่านก่อน-หลังของมิเตอร์หลายตัว (ระบุเป็น list) ในรอบบิลหนึ่งๆ
    คืนค่าเป็น dict {meter_id: {'Reading_before':.., 'Reading_after':..}} -- ไม่มี key ถ้ามิเตอร์นั้น
    ยังไม่เคยมีการบันทึกเลขอ่านในรอบบิลนี้เลย
    ใช้กับหน้า edit_meter ที่ต้องโชว์เลขอ่านของมิเตอร์น้ำ+ไฟ 2 ตัวพร้อมกันในรอบบิลที่เลือก
    """
    meter_ids = [m for m in meter_ids if m]
    if not meter_ids:
        return {}

    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        placeholders = ','.join('?' for _ in meter_ids)
        cursor.execute(
            f"""
            SELECT Meter_id, Reading_before, Reading_after
            FROM dbo.Contract_meter_reading_tr
            WHERE Billing_year_be = ? AND Billing_month = ? AND Meter_id IN ({placeholders})
            """,
            billing_year_be, billing_month, *meter_ids,
        )
        columns = [c[0] for c in cursor.description]
        rows = [dict(zip(columns, r)) for r in cursor.fetchall()]
        return {r['Meter_id']: r for r in rows}
    finally:
        conn.close()


def fetch_meters(type_filter=None, search=None):
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        sql = """
            SELECT
                m.Meter_id, m.SubArea_id, m.Meter_type_cd,
                m.Meter_no, m.Meter_no_status,
                STRING_AGG(c.Contract_code, ', ') AS bound_contracts
            FROM Contract_meter_ms m
            LEFT JOIN Contract_meter_tr t ON t.Meter_id = m.Meter_id
            LEFT JOIN Contract_hrd_tr c ON c.Contract_id = t.Contract_id
            WHERE 1=1
        """
        params = []
        if type_filter in ('8', '9'):
            sql += " AND m.Meter_type_cd = ?"
            params.append(int(type_filter))
        if search:
            sql += """ AND (
                m.SubArea_id LIKE ? OR
                m.Meter_no LIKE ? OR
                EXISTS (
                    SELECT 1 FROM Contract_meter_tr t2
                    JOIN Contract_hrd_tr c2 ON c2.Contract_id = t2.Contract_id
                    WHERE t2.Meter_id = m.Meter_id AND c2.Contract_code LIKE ?
                )
            )"""
            like = f"%{search}%"
            params += [like, like, like]
        sql += " GROUP BY m.Meter_id, m.SubArea_id, m.Meter_type_cd, m.Meter_no, m.Meter_no_status"
        sql += " ORDER BY m.Meter_id DESC"

        cursor.execute(sql, params)
        columns = [c[0] for c in cursor.description]
        rows = [dict(zip(columns, row)) for row in cursor.fetchall()]
        for r in rows:
            r['bound_contracts'] = (r['bound_contracts'] or '').split(', ') if r['bound_contracts'] else []
        return rows
    finally:
        conn.close()


def build_subareas_workbook(rows):
    """
    สร้าง Excel workbook (openpyxl) จากรายการ SubArea + มิเตอร์น้ำ/ไฟ ที่ได้จาก fetch_subareas()
    จัดฟอร์แมตหัวตาราง สี border ความกว้างคอลัมน์ และ auto-filter ให้พร้อมใช้งานทันที
    ใช้กับปุ่ม "ส่งออก Excel" บนหน้า dashboard -- คืนค่าเป็น Workbook object (ผู้เรียกเป็นคน save ต่อ)
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'มิเตอร์น้ำ-ไฟ'

    headers = [
        'เลขที่สัญญา (SubArea)', 'รหัสสัญญา', 'รหัสลูกหนี้', 'ชื่อลูกหนี้',
        'พื้นที่', 'สถานที่ตั้ง', 'เลขมิเตอร์น้ำ', 'เลขมิเตอร์ไฟฟ้า', 'สถานะ',
    ]

    header_fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    thin = Side(style='thin', color='D1D5DB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'

    for r in rows:
        status = 'มีสัญญา' if r.get('contract_code') else 'ยังไม่ผูกสัญญา'
        ws.append([
            r.get('SubArea_id') or '',
            r.get('contract_code') or '',
            r.get('customer_id') or '',
            r.get('customer_name') or '',
            r.get('Area_id') or '',
            r.get('SubArea_name') or '',
            r.get('water_meter_no') or '',
            r.get('electric_meter_no') or '',
            status,
        ])

    last_row = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical='center')

    widths = [20, 16, 14, 28, 16, 22, 16, 16, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"

    return wb


def fetch_dashboard_stats():
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT
                (SELECT COUNT(*) FROM Contract_location_subarea_ms) AS subarea_count,
                (SELECT COUNT(DISTINCT Contract_id) FROM Contract_meter_tr) AS contract_count,
                (SELECT COUNT(*) FROM Contract_meter_ms WHERE Meter_type_cd = 9) AS water_count,
                (SELECT COUNT(*) FROM Contract_meter_ms WHERE Meter_type_cd = 8) AS electric_count,
                (SELECT COUNT(*) FROM Contract_meter_ms WHERE Meter_no_status = N'รอ Gen เลข') AS pending_gen
        """)
        columns = [c[0] for c in cursor.description]
        row = cursor.fetchone()
        return dict(zip(columns, row)) if row else {}
    finally:
        conn.close()


def save_meter_reading(meter_id, billing_year_be, billing_month, reading_before, reading_after, user_id):
    """
    บันทึกเลขอ่านก่อน-หลังของมิเตอร์ 1 ตัว ในรอบบิลหนึ่ง (upsert)
    ผ่าน sp_Contract_Meter_Reading_Save -- คืนค่า Reading_id ที่บันทึก
    """
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        cursor.execute(
            """
            EXEC dbo.sp_Contract_Meter_Reading_Save
                @Meter_id        = ?,
                @Billing_year_be = ?,
                @Billing_month   = ?,
                @Reading_before  = ?,
                @Reading_after   = ?,
                @UserId          = ?
            """,
            meter_id, billing_year_be, billing_month, reading_before, reading_after, user_id,
        )
        row = cursor.fetchone()
        conn.commit()
        return int(row[0]) if row else None
    finally:
        conn.close()
# meters/views/download.py
import datetime
import io
from urllib.parse import quote

from django.http import HttpResponse
from django.shortcuts import render, redirect
from openpyxl import Workbook
from openpyxl.styles import Font

from .. import services
from .dashboard import THAI_MONTHS, _current_year_be, _safe_int


WATER_HEADERS = [
    'ปีที่สร้างสัญญา', 'รหัสลูกหนี้', 'ชื่อลูกหนี้', 'เลข13หลัก', 'เลขที่สัญญา', 'รหัสสัญญา',
    'ลำดับพื้นที่ย่อย', 'สถานที่ตั้ง', 'Area_id', 'พื้นที่ตามที่ตั้งพื้นที่เช่า', 'SubArea_id',
    'พื้นที่ย่อยตามที่ตั้งพื้นที่เช่า', 'สถานะสัญญา',
    'ยืนยันพื้นที่ที่คิดค่าน้ำ โดยใส่ 1 = มี, 0 = ไม่มี', 'ประเภทค่าใช้จ่าย',
    'ลำดับมิเตอร์WW_seq หากมี > 1เลข ให้ใส่เลขลำดับ เช่น 1, 2, 3…', 'เลขประจำเครื่องวัดน้ำประปา',
    'เลขที่อ่านครั้งหลัง', 'เลขที่อ่านครั้งก่อน', 'วันที่อ่านเลขมิเตอร์น้ำ', 'จำนวนเงินที่ต้องชำระ',
    'Document Date', 'Posting Date',
    'Document Type ส่วนพัฒนากายภาพ = DR, ส่วนพัฒนาความยั่งยืน = 61',
    'Base Line Date', 'วันที่ Upload ข้อมูล', 'ผู้Uploadข้อมูล', 'หมายเหตุ',
]

ELECTRIC_HEADERS = [
    'ปีที่สร้างสัญญา', 'รหัสลูกหนี้', 'ชื่อลูกหนี้', 'เลข13หลัก', 'เลขที่สัญญา', 'รหัสสัญญา',
    'ลำดับพื้นที่ย่อย', 'Location_id', 'สถานที่ตั้ง', 'Area_id',
    'พื้นที่ตามที่ตั้งพื้นที่เช่า', 'SubArea_id', 'พื้นที่ย่อยตามที่ตั้งพื้นที่เช่า', 'สถานะสัญญา',
    'ยืนยันพื้นที่ที่คิดค่าไฟ โดยใส่ 1 = มี, 0 = ไม่มี', 'ประเภทค่าใช้จ่าย',
    'ลำดับมิเตอร์EE_seq หากมี > 1เลข ให้ใส่เลขลำดับ เช่น 1, 2, 3…', 'เลขประจำเครื่องวัดไฟฟ้า',
    'กรอกระบบไฟฟ้า 2 ประเภท คือ ระบบไฟฟ้า 1 เฟส, ระบบไฟฟ้า 3 เฟส',
    'เลขที่อ่านครั้งหลัง', 'เลขที่อ่านครั้งก่อน', 'วันที่อ่านเลขมิเตอร์ไฟฟ้า', 'จำนวนเงินที่ต้องชำระ',
    'Document Date', 'Posting Date',
    'Document Type ส่วนพัฒนากายภาพ = DR, ส่วนพัฒนาความยั่งยืน = 61',
    'Base Line Date', 'วันที่ Upload ข้อมูล', 'ผู้Uploadข้อมูล', 'หมายเหตุ',
]


def _combined_key(r):
    """
    key รวมสำหรับ merge แถวน้ำ+ไฟที่เป็นสัญญา/พื้นที่ย่อยเดียวกันให้อยู่แถวเดียวกันในหน้าจอ
    Contract_id + SubArea_id + _slot (ไม่รวม Meter_id เพราะต้องการรวมน้ำ+ไฟของพื้นที่ย่อย
    เดียวกันเข้าด้วยกัน ไม่ใช่แยกตามมิเตอร์เหมือน row_key เดิม)

    _slot คือ "มิเตอร์ตัวที่ n ของประเภทนั้นในสัญญา+พื้นที่ย่อยนี้" (นับจาก 1 ตามลำดับที่ query คืนมา
    ซึ่งเรียงด้วย Meter_id) -- เดิม key ไม่มี slot พื้นที่ที่มีมิเตอร์น้ำ/ไฟประเภทเดียวกันหลายตัว
    (เช่น A10 มีน้ำ 1014 กับ 1038) จะถูกรวมเป็นแถวเดียวแล้วตัวหลังเขียนทับตัวแรกบนหน้าจอ ทั้งที่ตอน
    export ไฟล์ออกครบทุกตัว ตอนนี้ตัวที่ n ของน้ำจับคู่กับตัวที่ n ของไฟ และมิเตอร์ที่เกินมา
    ได้แถวของตัวเอง ทุกตัวจึงเห็นและเลือกได้
    """
    return f"{r['Contract_id']}|{r['SubArea_id']}|{r['_slot']}"


def _merge_water_electric(water_rows, electric_rows):
    """
    รวมแถวน้ำ+ไฟที่เป็นสัญญา/พื้นที่ย่อย/slot เดียวกันเข้าด้วยกัน เพื่อโชว์เป็นแถวเดียวในหน้าจอเลือก
    (เหมือนการ์ดรวมน้ำ-ไฟใน edit_meter.html) -- แต่ตอน export Excel ยังคงแยก 2 ชีทเหมือนเดิม
    """
    combined = {}

    for kind, rows in (('water', water_rows), ('electric', electric_rows)):
        for r in rows:
            key = _combined_key(r)
            entry = combined.setdefault(key, {
                'row_key': key,
                'Contract_code': r.get('Contract_code'),
                'CompanyName': r.get('CompanyName'),
                'SubArea_id': r.get('SubArea_id'),
                'slot': r['_slot'],
                'water': None,
                'electric': None,
            })
            entry[kind] = r

    return sorted(
        combined.values(),
        key=lambda x: (x['Contract_code'] or '', x['SubArea_id'] or '', x['slot']),
    )


def _dedupe_meter_rows(rows):
    """
    ตัดแถวซ้ำก่อนเขียนลง Excel -- query จาก services.fetch_contracts_for_download บางครั้งคืน
    แถวซ้ำสำหรับมิเตอร์เดียวกัน (ข้อมูลที่จะแสดงในไฟล์เหมือนกันทุกอย่าง) เข้าใจว่าเกิดจาก JOIN
    ตารางงวดเดิม (Contract_Installment_tr_dt) ที่ผูกด้วยสัญญา+ประเภทค่าใช้จ่าย+เดือน ไม่ผูกกับมิเตอร์
    จึง fan-out ซ้ำ ควรตามไปแก้ที่ต้นตอใน SQL query ด้วย

    ตั้งใจ dedupe ด้วยค่าที่ "แสดงผลจริง" ในไฟล์ (Meter_no ที่โชว์ + เลขอ่าน/จำนวนเงิน) แทนที่จะ
    ใช้ Meter_id ภายใน เพราะพบว่าบางครั้งแถวซ้ำมี Meter_id คนละค่ากัน (2 เรคคอร์ดในฐานข้อมูลที่ชี้
    ไปมิเตอร์คนละตัวแต่ Meter_no ที่แสดงเหมือนกัน) -- dedupe แบบเดิมที่ใช้ Meter_id จึงจับไม่ได้
    """
    seen = set()
    result = []
    for r in rows:
        key = (
            r['Contract_id'],
            r['SubArea_id'],
            r.get('Meter_no'),
            r.get('Contract_read_number_before'),
            r.get('Contract_read_number_after'),
            r.get('Contract_effective_date'),
            r.get('Contract_Installment_amt'),
        )
        if key in seen:
            continue
        seen.add(key)
        result.append(r)
    return result


def _prepare_rows(meter_type_cd, year_be, month):
    """
    ดึงข้อมูล -> ตัดแถวซ้ำ -> ใส่ _slot และ _seq ให้ "ทั้งชุด" ก่อนจะมีการกรองตามที่ผู้ใช้เลือก

    ต้องคำนวณก่อนกรองเสมอ ทั้งสองค่าเป็นของตำแหน่งในชุดเต็ม:
      - _slot ใช้ประกอบ row_key ที่หน้าจอกับตอน export ต้องได้ค่าเดียวกัน
      - _seq คือลำดับมิเตอร์ WW_seq/EE_seq ต่อสัญญา (สมมติฐานชั่วคราว ดูคอมเมนต์ใน services/download.py)
        เดิมนับหลังกรอง ถ้าติ๊กเลือกเฉพาะมิเตอร์ตัวที่ 2 ของสัญญา มันจะกลายเป็น seq = 1 ในไฟล์
    """
    rows = _dedupe_meter_rows(services.fetch_contracts_for_download(meter_type_cd, year_be, month))
    slot_counters = {}
    seq_counters = {}
    for r in rows:
        slot_key = (r['Contract_id'], r['SubArea_id'])
        slot_counters[slot_key] = slot_counters.get(slot_key, 0) + 1
        r['_slot'] = slot_counters[slot_key]
        seq_counters[r['Contract_id']] = seq_counters.get(r['Contract_id'], 0) + 1
        r['_seq'] = seq_counters[r['Contract_id']]
    return rows


def _append_row(ws, values):
    """
    เขียนแถวลงชีท โดยกันข้อความที่ขึ้นต้นด้วย '=' ไม่ให้ openpyxl ตีเป็นสูตร (เช่นชื่อลูกหนี้หรือ
    รหัสที่มาจากข้อมูลภายนอก) -- ถ้าเป็นสูตร Excel จะคำนวณตอนเปิดไฟล์แทนที่จะแสดงข้อความตามที่เก็บไว้
    """
    ws.append(values)
    for cell in ws[ws.max_row]:
        if cell.data_type == 'f':
            cell.data_type = 's'


def _write_water_sheet(ws, rows):
    ws.append(WATER_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in rows:
        _append_row(ws, [
            r['Contract_year'], r['Customer_id'], r['CompanyName'], r['TaxNumber'],
            r['Contract_id'], r['Contract_code'], r['Contract_Location_seq'], r['Location_name'],
            r['Area_id'], r['Area_name'], r['SubArea_id'], r['Contract_area_desc'],
            r['Status_contract_Desc'], 1, 'ค่าน้ำประปา', r['_seq'], r['Meter_no'],
            r['Contract_read_number_after'], r['Contract_read_number_before'],
            r['Contract_effective_date'], r['Contract_Installment_amt'],
            None, None, 'DR', r['Contract_duedate'], None, None, None,
        ])


def _write_electric_sheet(ws, rows):
    ws.append(ELECTRIC_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in rows:
        _append_row(ws, [
            r['Contract_year'], r['Customer_id'], r['CompanyName'], r['TaxNumber'],
            r['Contract_id'], r['Contract_code'], r['Contract_Location_seq'], r['Location_id'],
            r['Location_name'], r['Area_id'], r['Area_name'], r['SubArea_id'],
            r['Contract_area_desc'], r['Status_contract_Desc'], 1, 'ค่าไฟฟ้า', r['_seq'],
            r['Meter_no'], r['Phase_type'], r['Contract_read_number_after'],
            r['Contract_read_number_before'], r['Contract_effective_date'],
            r['Contract_Installment_amt'], None, None, 'DR', r['Contract_duedate'],
            None, None, None,
        ])


def _year_options():
    return [_current_year_be() - 1, _current_year_be(), _current_year_be() + 1]


def download_form(request):
    """
    หน้าเลือกเดือน-ปี + ค้นหา -- GET ธรรมดา (ไม่มี month/year) โชว์แค่ฟอร์มเปล่า
    ถ้ามี ?month=&year= (กดปุ่ม "ค้นหา") จะ query ข้อมูลจริงมาโชว์เป็นตารางเดียวที่รวมน้ำ+ไฟ
    ต่อสัญญา/พื้นที่ย่อยในหน้าเดียวกันเลย (เหมือนการ์ดรวมใน edit_meter.html) พร้อม checkbox
    ให้ติ๊กเลือกก่อนค่อยกดดาวน์โหลด Excel จริง (submit ไป download_excel แยกต่างหาก)
    """
    context = {
        'thai_months': THAI_MONTHS,
        'billing_month': datetime.date.today().month,
        'billing_year': _current_year_be(),
        'year_options': _year_options(),
        'searched': False,
    }

    if 'month' in request.GET or 'year' in request.GET:
        # เดิมสร้าง errors = [] ส่งเข้า _safe_int แล้วไม่เคยใส่ใน context
        # ส่ง ?year=abc เข้ามาจะ fallback เป็นปีปัจจุบันแบบเงียบ ผู้ใช้ได้ข้อมูลของรอบบิล
        # ที่ไม่ได้ขอโดยไม่รู้ตัว -- ตอนนี้ตรวจช่วงด้วยและส่งข้อความไปแสดงบนหน้าจอ
        # (เทียบ `is None` ไม่ใช้ `or` เพราะ 0 เป็น falsy จะถูกกลืนเป็น fallback เงียบๆ)
        period_errors = []
        year_be = _safe_int(request.GET.get('year'), 'ปี', period_errors,
                            min_value=2400, max_value=2700)
        if year_be is None:
            year_be = _current_year_be()
        month = _safe_int(request.GET.get('month'), 'เดือน', period_errors,
                          min_value=1, max_value=12)
        if month is None:
            month = datetime.date.today().month

        context.update({
            'billing_month': month,
            'billing_year': year_be,
            'period_errors': period_errors,
        })

        # ถ้ารอบบิลที่ส่งมาไม่ถูกต้อง ไม่ query ข้อมูลเลย -- ไม่ควรโชว์ตารางของรอบบิลอื่น
        # ให้ผู้ใช้ติ๊กเลือกแล้วดาวน์โหลดไป ทั้งที่เขาคิดว่าเป็นรอบบิลที่ตัวเองกรอก
        if not period_errors:
            water_rows = _prepare_rows(9, year_be, month)
            electric_rows = _prepare_rows(8, year_be, month)

            context.update({
                'searched': True,
                'combined_rows': _merge_water_electric(water_rows, electric_rows),
            })

    return render(request, 'meters/download_form.html', context)


def _form_with_errors(request, month, year_be, errors):
    """กลับไปหน้าเลือกพร้อมข้อความ แทนที่จะส่งไฟล์เปล่า/ไฟล์ของรอบบิลที่ผู้ใช้ไม่ได้ขอ"""
    return render(request, 'meters/download_form.html', {
        'thai_months': THAI_MONTHS,
        'billing_month': month if month is not None else datetime.date.today().month,
        'billing_year': year_be if year_be is not None else _current_year_be(),
        'year_options': _year_options(),
        'searched': False,
        'period_errors': errors,
    })


def download_excel(request):
    """
    สร้างไฟล์ Excel จากแถวที่ผู้ใช้ติ๊กเลือกในหน้าเว็บเท่านั้น (POST) -- query ข้อมูลเต็มใหม่อีกรอบ
    ด้วย month/year เดิม แล้วกรองเอาเฉพาะแถวที่ key รวม (Contract_id|SubArea_id|slot) ตรงกับที่ติ๊กมา
    (ไม่เชื่อข้อมูลที่ POST มาตรงๆ เพราะ client ส่งอะไรมาก็ได้ -- query ใหม่จาก DB เสมอ ปลอดภัยกว่า)

    หน้าจอเลือกข้อมูลรวมน้ำ+ไฟเป็นแถวเดียว แต่ไฟล์ Excel ที่ export ออกมายังคงแยก 2 ชีท
    (น้ำ/ไฟ) เหมือนเดิม เพราะโครงสร้างคอลัมน์ตั้งต้นของแต่ละประเภทไม่เหมือนกัน
    """
    # ต้องเป็น POST เท่านั้น -- เดิมเข้าด้วย GET ได้ แล้ว request.POST ว่าง ทำให้
    # selected_rows ว่าง ได้ไฟล์ Excel ที่มีแต่หัวตาราง ผู้ใช้เข้าใจว่า "ไม่มีข้อมูล"
    # ทั้งที่จริงคือเข้ามาผิดทาง จึงเด้งกลับไปหน้าเลือกรอบบิลแทน
    if request.method != 'POST':
        return redirect('download_form')

    period_errors = []
    year_be = _safe_int(request.POST.get('year'), 'ปี', period_errors,
                        min_value=2400, max_value=2700)
    month = _safe_int(request.POST.get('month'), 'เดือน', period_errors,
                      min_value=1, max_value=12)

    selected_keys = set(request.POST.getlist('selected_rows'))

    # รอบบิลเพี้ยนหรือไม่ได้ติ๊กอะไรมา -- กลับไปหน้าเลือกพร้อมข้อความ ดีกว่าส่งไฟล์เปล่า
    # หรือไฟล์ของรอบบิลที่ผู้ใช้ไม่ได้ขอ ให้ผู้ใช้เอาไปตั้งหนี้ต่อ
    if period_errors or not selected_keys:
        if not selected_keys:
            period_errors.append('ยังไม่ได้เลือกรายการที่ต้องการดาวน์โหลด -- กรุณาติ๊กอย่างน้อย 1 รายการ')
        return _form_with_errors(request, month, year_be, period_errors)

    water_rows = [r for r in _prepare_rows(9, year_be, month) if _combined_key(r) in selected_keys]
    electric_rows = [r for r in _prepare_rows(8, year_be, month) if _combined_key(r) in selected_keys]

    # ติ๊กมาแต่ไม่ตรงกับรายการไหนเลย (เช่น หน้าค้างจากเวอร์ชันก่อนที่ key ยังไม่มี slot
    # หรือข้อมูลเปลี่ยนไประหว่างนั้น) -- บอกผู้ใช้ให้ค้นหาใหม่ ไม่ส่งไฟล์ที่มีแต่หัวตาราง
    if not water_rows and not electric_rows:
        return _form_with_errors(
            request, month, year_be,
            ['ไม่พบรายการที่เลือกในรอบบิลนี้ -- ข้อมูลอาจเปลี่ยนไปแล้ว กรุณากดค้นหาใหม่แล้วเลือกอีกครั้ง'],
        )

    wb = Workbook()
    ws1 = wb.active
    ws1.title = f"สัญญาปี{year_be}_น้ำประปา"[:31]
    _write_water_sheet(ws1, water_rows)

    ws2 = wb.create_sheet(title=f"สัญญาปี{year_be}_ไฟฟ้า"[:31])
    _write_electric_sheet(ws2, electric_rows)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"สัญญาน้ำ-ไฟ_{month:02d}-{year_be}.xlsx"
    # ชื่อไฟล์มีอักขระไทย -- HTTP header ต้องเป็น ASCII เท่านั้น ถ้าใส่ภาษาไทยดิบๆ ลงไป
    # เบราว์เซอร์บางตัว (เช่น Chrome) จะอ่าน header ไม่ออกแล้ว fallback ไปใช้ชื่อ default
    # ("ดาวน์โหลด.xlsx" ตาม locale) แทน จึงต้อง encode ตาม RFC 5987 (filename*=UTF-8''...)
    # ควบคู่กับ filename="..." แบบ ASCII fallback (เผื่อเบราว์เซอร์เก่าที่ไม่รองรับ filename*)
    ascii_fallback = f"download_{month:02d}-{year_be}.xlsx"
    encoded_filename = quote(filename)
    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename="{ascii_fallback}"; '
        f"filename*=UTF-8''{encoded_filename}"
    )
    return response
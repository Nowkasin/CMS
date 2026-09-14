# meters/services/excel_export.py
"""
สร้างไฟล์ Excel สำหรับปุ่ม "ส่งออก Excel" บนหน้า dashboard
"""


def build_subareas_workbook(rows):
    """
    สร้าง Excel workbook (openpyxl) จากรายการ SubArea + มิเตอร์น้ำ/ไฟ ที่ได้จาก fetch_subareas()
    จัดฟอร์แมตหัวตาราง สี border ความกว้างคอลัมน์ และ auto-filter ให้พร้อมใช้งานทันที
    คืนค่าเป็น Workbook object (ผู้เรียกเป็นคน save ต่อ)
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'มิเตอร์น้ำ-ไฟ'

    headers = [
        'เลขที่สัญญา (SubArea)', 'รหัสสัญญา', 'รหัสลูกหนี้', 'ชื่อลูกหนี้',
        'พื้นที่', 'สถานที่ตั้ง', 'เลขมิเตอร์น้ำ', 'เลขมิเตอร์ไฟฟ้า', 'สถานะ',
    ]

    header_fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    thin = Side(style='thin', color='D1D5DB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'

    for r in rows:
        status = 'มีสัญญา' if r.get('contract_code') else 'ยังไม่ผูกสัญญา'
        ws.append([
            r.get('SubArea_id') or '',
            r.get('contract_code') or '',
            r.get('customer_id') or '',
            r.get('customer_name') or '',
            r.get('Area_id') or '',
            r.get('SubArea_name') or '',
            r.get('water_meter_no') or '',
            r.get('electric_meter_no') or '',
            status,
        ])

    last_row = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical='center')

    widths = [20, 16, 14, 28, 16, 22, 16, 16, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"

    return wb